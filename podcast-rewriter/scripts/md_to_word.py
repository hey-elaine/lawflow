"""Convert podcast markdown scripts to Word documents.

Usage: python md_to_word.py <input.md> <output.docx>

Body text: 宋体 11pt, 3x line spacing. Headings: blue (RGB 31,73,125).
"""
import re
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def convert(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    # Normal style: 宋体 11pt 3x line spacing
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 3.0

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped == '---':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run('— — — — —')
            run.font.color.rgb = RGBColor(180, 180, 180)
            continue

        heading_match = re.match(r'^(#{1,6})\s+(.+)$', stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            h = doc.add_heading(text, level=min(level, 3))
            for run in h.runs:
                run.font.color.rgb = RGBColor(31, 73, 125)
            continue

        # Strip markdown formatting
        text = stripped
        text = re.sub(r'\*{1,3}(.+?)\*{1,3}', r'\1', text)
        text = re.sub(r'`(.+?)`', r'\1', text)
        text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', text)

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(22)
        p.paragraph_format.line_spacing = 3.0
        run = p.add_run(text)
        run.font.name = '宋体'
        run.font.size = Pt(11)

    doc.save(docx_path)
    print(f'Done: {docx_path}')


if __name__ == '__main__':
    if len(sys.argv) >= 3:
        convert(sys.argv[1], sys.argv[2])
    else:
        print('Usage: python md_to_word.py <input.md> <output.docx>')
